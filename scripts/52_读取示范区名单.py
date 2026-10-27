# -*- coding: utf-8 -*-
"""52_读取示范区名单.py：提取 2023 年示范县名单 → 存 UTF-8 文本"""
from pathlib import Path
from docx import Document

d = Document(r"E:\GLM\非遗论文\01_选题与文献\2023年传统村落集中连片保护利用示范县（市、区）名单.docx")
lines = []
for p in d.paragraphs:
    t = p.text.strip()
    if t:
        lines.append(t)
for ti, table in enumerate(d.tables):
    lines.append(f"--- 表{ti+1} ---")
    for row in table.rows:
        lines.append(" ｜ ".join(c.text.strip() for c in row.cells))

out = Path(r"E:\GLM\非遗论文\01_选题与文献\示范区名单_2023_提取.txt")
out.write_text("\n".join(lines), encoding="utf-8")
print(f"已保存：{out}（{len(lines)} 段）")
