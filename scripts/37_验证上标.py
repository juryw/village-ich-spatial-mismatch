# -*- coding: utf-8 -*-
"""37_验证上标.py：抽查 [9] 等引文编号是否为真正的上标 run"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from docx import Document

d = Document(r"E:\GLM\非遗论文\05_稿件\论文初稿_v1.docx")
sup_runs = []
for p in d.paragraphs:
    for r in p.runs:
        if r.font.superscript and r.text.strip():
            sup_runs.append(r.text)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if r.font.superscript and r.text.strip():
                        sup_runs.append(r.text + "（表内）")
print(f"上标 run 总数：{len(sup_runs)}")
print("样例:", sup_runs[:8], "...", sup_runs[-4:])
