# -*- coding: utf-8 -*-
"""47_验证图片嵌入.py"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from docx import Document

d = Document(r"E:\GLM\非遗论文\05_稿件\论文初稿_v1.docx")
print("内嵌图片数:", len(d.inline_shapes), "（应为 5）")
print("表格数:", len(d.tables), "（应为 5：数据来源+表2+表3+表4+... 实际核对）")
caps = [p.text for p in d.paragraphs if p.text.strip().startswith("图 ") and "——" not in p.text]
for c in caps:
    print("图题:", c)
