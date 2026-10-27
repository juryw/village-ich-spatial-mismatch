# -*- coding: utf-8 -*-
"""35_验证Word.py"""
from docx import Document

d = Document(r"E:\GLM\非遗论文\05_稿件\论文初稿_v1.docx")
print("段落数:", len(d.paragraphs), "| 表格数:", len(d.tables))
print("题目:", d.paragraphs[0].text[:45])
refs = [p.text for p in d.paragraphs if p.text.strip().startswith("[36]")]
print("参考文献[36]:", (refs[0][:60] + "...") if refs else "缺失!")
tables_ok = all(len(t.rows) >= 2 for t in d.tables)
print("表格完整性:", "OK" if tables_ok else "异常")
