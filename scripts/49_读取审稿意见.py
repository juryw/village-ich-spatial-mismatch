# -*- coding: utf-8 -*-
"""49_读取审稿意见.py：提取 DeepSeek 审稿意见 docx 的全部文字 → 存为 UTF-8 文本"""
from pathlib import Path
from docx import Document

d = Document(r"C:\Users\asus\Documents\deepseek-harness\default-workspace\审稿意见_中国传统村落与非遗空间错位_20261028.docx")
lines = []
for p in d.paragraphs:
    t = p.text.strip()
    if t:
        lines.append(t)
for ti, table in enumerate(d.tables):
    lines.append(f"--- 表{ti+1} ---")
    for row in table.rows:
        lines.append(" | ".join(c.text.strip() for c in row.cells))

out = Path(r"E:\GLM\非遗论文\01_选题与文献\DeepSeek审稿意见_全文提取.txt")
out.write_text("\n".join(lines), encoding="utf-8")
print(f"已保存：{out}（{len(lines)} 段）")
